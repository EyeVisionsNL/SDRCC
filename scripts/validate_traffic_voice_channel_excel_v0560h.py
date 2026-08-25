#!/usr/bin/env python3
from copy import deepcopy
from pathlib import Path
from io import BytesIO
import sys
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from core import traffic_voice

def check(ok, message):
    if not ok:
        print(f"FAIL: {message}")
        raise SystemExit(1)
    print(f"PASS: {message}")

check((ROOT / "VERSION").read_text().strip() == "0.56.0h", "release version is 0.56.0h")
check("openpyxl" in (ROOT / "requirements.txt").read_text().lower(), "openpyxl dependency is explicit")
check(hasattr(traffic_voice, "export_channel_list_xlsx"), "Traffic Voice exposes Excel export")
check(hasattr(traffic_voice, "import_channel_list_xlsx"), "Traffic Voice exposes bounded Excel import")

app = (ROOT / "dashboard/app.py").read_text()
html = (ROOT / "dashboard/templates/index.html").read_text()
js = (ROOT / "dashboard/static/js/traffic_voice.js").read_text()
core_text = (ROOT / "core/traffic_voice.py").read_text()

check("/api/traffic-voice/channel-list.xlsx" in app, "bounded Excel export route is present")
check("/api/traffic-voice/channel-list/import" in app, "bounded Excel import route is present")
check('id="traffic-voice-channel-list-load"' in html, "Load Excel control is present")
check('id="traffic-voice-channel-list-export"' in html, "Export Excel control is present")
check("FormData" in js and "channelListImportEndpoint" in js, "browser multipart import is wired")
check("config_core.save_traffic_voice(candidate)" in core_text, "existing Traffic Voice YAML remains write authority")

raw = traffic_voice.config_core.load_traffic_voice()
blob = traffic_voice.export_channel_list_xlsx(raw)
book = load_workbook(BytesIO(blob), read_only=True, data_only=True)
check("Channels" in book.sheetnames and "Instructions" in book.sheetnames, "export contains Channels and Instructions sheets")
headers = [cell.value for cell in next(book["Channels"].iter_rows(min_row=1, max_row=1))]
check(tuple(headers[:7]) == traffic_voice.CHANNEL_LIST_HEADERS, "export headers match import contract")

stored = deepcopy(raw)
writes = []
orig_load = traffic_voice.config_core.load_traffic_voice
orig_save = traffic_voice.config_core.save_traffic_voice
try:
    traffic_voice.config_core.load_traffic_voice = lambda: deepcopy(stored)
    traffic_voice.config_core.save_traffic_voice = lambda candidate: writes.append(deepcopy(candidate))
    result = traffic_voice.import_channel_list_xlsx(blob, source_name="roundtrip.xlsx")
    check(result["marine_channels"] > 0 and result["aviation_channels"] > 0, "round-trip retains both channel banks")
    check(len(writes) == 1, "valid workbook performs one bounded YAML write")
    before_invalid = len(writes)
    try:
        traffic_voice.import_channel_list_xlsx(b"not an xlsx", source_name="broken.xlsx")
    except ValueError:
        pass
    else:
        raise AssertionError("invalid workbook was accepted")
    check(len(writes) == before_invalid, "invalid workbook performs no YAML write")
finally:
    traffic_voice.config_core.load_traffic_voice = orig_load
    traffic_voice.config_core.save_traffic_voice = orig_save

check((ROOT / "docs/traffic-voice-channel-list-excel-v0560h.md").exists(), "v0.56.0h documentation is present")
print("VALIDATION PASS: SDRCC v0.56.0h Traffic Voice Excel Channel Lists")
